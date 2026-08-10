"""Inventory CALCE archive extraction without transforming it."""
from __future__ import annotations

import argparse
import json
from collections import Counter
from pathlib import Path

import openpyxl


def _text_preview(path: Path) -> list[str]:
    with path.open("r", encoding="utf-8", errors="replace") as handle:
        return [line.rstrip("\r\n") for _, line in zip(range(3), handle)]


def _sheet_header(sheet) -> list[str]:
    iter_rows = getattr(sheet, "iter_rows", None)
    if iter_rows is None:
        return []
    return ["" if value is None else str(value) for value in next(iter_rows(max_row=1, values_only=True), ())]


def _workbook_summary(path: Path) -> dict[str, list[str]]:
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return {sheet.title: _sheet_header(sheet) for sheet in book._sheets}
    finally:
        book.close()


def inspect_directory(root: Path) -> dict:
    files = sorted(path for path in root.rglob("*") if path.is_file())
    extensions = Counter((path.suffix.lower() or "[no extension]") for path in files)
    return {
        "root": str(root),
        "file_count": len(files),
        "extensions": dict(sorted(extensions.items())),
        "files": [str(path.relative_to(root)) for path in files],
        "text_previews": {str(path.relative_to(root)): _text_preview(path) for path in files if path.suffix.lower() in {".txt", ".csv"}},
        "workbooks": {str(path.relative_to(root)): _workbook_summary(path) for path in files if path.suffix.lower() == ".xlsx"},
    }


def write_inventory(root: Path, output: Path) -> dict:
    summary = inspect_directory(root)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return summary


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    summary = write_inventory(args.input_dir, args.output)
    print(json.dumps({"file_count": summary["file_count"], "extensions": summary["extensions"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
